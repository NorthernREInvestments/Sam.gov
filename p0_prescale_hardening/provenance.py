"""P0-3 — SOURCE_PROVENANCE for every material line (PDF page / spreadsheet sheet-row-cell)."""

from __future__ import annotations

import hashlib
import re
from functools import lru_cache
from pathlib import Path
from typing import Any

from p0_prescale_hardening.models import NO_PROVENANCE, NOT_QUOTE_READY

try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover
    load_workbook = None  # type: ignore


def _fname(path: Any) -> str | None:
    if not path:
        return None
    return str(path).replace("\\", "/").rsplit("/", 1)[-1]


def _doc_id(path: Any) -> str | None:
    if not path:
        return None
    s = str(path)
    return "DOC-" + hashlib.sha1(s.encode()).hexdigest()[:12]


@lru_cache(maxsize=32)
def _xlsx_index(path_str: str) -> dict[str, dict[str, Any]]:
    """Map normalized token → first sheet/row/cell hit."""
    out: dict[str, dict[str, Any]] = {}
    if load_workbook is None:
        return out
    p = Path(path_str)
    if not p.exists():
        return out
    try:
        wb = load_workbook(p, read_only=True, data_only=True)
    except Exception:
        return out
    try:
        for sheet in wb.worksheets:
            for row_i, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                for col_i, val in enumerate(row, start=1):
                    if val is None:
                        continue
                    raw = str(val).strip()
                    if len(raw) < 3:
                        continue
                    key = re.sub(r"\s+", " ", raw).upper()
                    if key not in out:
                        from openpyxl.utils import get_column_letter

                        cell = f"{get_column_letter(col_i)}{row_i}"
                        out[key] = {
                            "sheet": sheet.title,
                            "row": row_i,
                            "column": col_i,
                            "cell": cell,
                            "raw_value": raw[:200],
                        }
    finally:
        wb.close()
    return out


def _lookup_xlsx(path: str, *tokens: Any) -> dict[str, Any] | None:
    idx = _xlsx_index(path)
    for t in tokens:
        if not t:
            continue
        key = re.sub(r"\s+", " ", str(t)).strip().upper()
        if key in idx:
            return idx[key]
        # partial: token contained in cell or cell contained in token
        for k, v in idx.items():
            if key in k or k in key:
                if len(key) >= 4 and len(k) >= 4:
                    return v
    return None


@lru_cache(maxsize=16)
def _pdf_pages(path_str: str) -> list[str]:
    try:
        from document_quality import extract_pdf_text_by_page

        pages = extract_pdf_text_by_page(path_str, max_pages=80)
        if pages:
            return list(pages)
    except Exception:
        pass
    try:
        from document_quality import extract_pdf_text

        text = extract_pdf_text(path_str, max_pages=40) or ""
        # crude single-blob fallback — page unknown
        return [text] if text else []
    except Exception:
        return []


def _lookup_pdf_page(path: str, *tokens: Any) -> int | None:
    pages = _pdf_pages(path)
    for t in tokens:
        if not t:
            continue
        needle = str(t).strip()
        if len(needle) < 4:
            continue
        for i, page_text in enumerate(pages, start=1):
            if needle.upper() in (page_text or "").upper():
                return i
    return None


def build_source_provenance(line: dict[str, Any]) -> dict[str, Any]:
    """Required SOURCE_PROVENANCE object."""
    src_path = line.get("source_document") or line.get("source_path")
    fname = _fname(src_path)
    doc_id = line.get("document_id") or _doc_id(src_path)
    kind = str(line.get("source_kind") or "").lower()
    if not kind and fname:
        kind = Path(fname).suffix.lstrip(".").lower()

    page = line.get("page")
    sheet = line.get("sheet")
    row = line.get("row")
    cell = line.get("cell") or line.get("cell_range")
    section = line.get("section")
    raw = line.get("raw_source_text") or line.get("description") or ""
    contributing = list(line.get("contributing_cells") or [])

    # Enrich spreadsheet
    if src_path and kind in {"xlsx", "xls", "csv"} and (sheet is None or row is None or cell is None):
        hit = _lookup_xlsx(
            str(src_path),
            line.get("mpn"),
            line.get("model"),
            line.get("part_number"),
            line.get("clin"),
            (line.get("description") or "")[:40],
        )
        if hit:
            sheet = sheet or hit.get("sheet")
            row = row or hit.get("row")
            cell = cell or hit.get("cell")
            if not contributing:
                contributing = [
                    {
                        "sheet": hit.get("sheet"),
                        "row": hit.get("row"),
                        "cell": hit.get("cell"),
                        "raw_value": hit.get("raw_value"),
                    }
                ]
            if not line.get("raw_source_text"):
                raw = hit.get("raw_value") or raw

    # Enrich PDF page
    if src_path and kind == "pdf" and page is None:
        page = _lookup_pdf_page(
            str(src_path),
            line.get("mpn"),
            line.get("model"),
            (line.get("description") or "")[:48],
        )

    prov = {
        "document_id": doc_id,
        "filename": fname,
        "source_kind": kind or None,
        "page": page,
        "section": section,
        "sheet": sheet,
        "row": row,
        "column": line.get("column"),
        "cell": cell,
        "cell_range": line.get("cell_range") or cell,
        "contributing_cells": contributing,
        "raw_source_text": (raw or "")[:400],
        "source_path": str(src_path) if src_path else None,
    }

    complete = bool(fname) and bool(raw)
    if kind in {"xlsx", "xls", "csv"}:
        complete = complete and sheet is not None and (row is not None or cell is not None)
    elif kind == "pdf":
        complete = complete and page is not None
    else:
        # unknown kind — require at least filename + locator of some kind
        complete = complete and (page is not None or sheet is not None or cell is not None)

    prov["complete"] = complete
    prov["status"] = "OK" if complete else NO_PROVENANCE
    return prov


def attach_provenance(line: dict[str, Any]) -> dict[str, Any]:
    out = dict(line)
    prov = build_source_provenance(line)
    out["SOURCE_PROVENANCE"] = prov
    if not prov.get("complete"):
        out["quote_ready_gate"] = NOT_QUOTE_READY
        out["quote_ready_blockers"] = list(set((out.get("quote_ready_blockers") or []) + [NO_PROVENANCE]))
    return out


def is_quote_ready_provenance(line: dict[str, Any]) -> bool:
    prov = line.get("SOURCE_PROVENANCE") or build_source_provenance(line)
    return bool(prov.get("complete"))
