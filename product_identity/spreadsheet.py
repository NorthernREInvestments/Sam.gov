"""Structured spreadsheet extraction — preserve columns, do not flatten."""

from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Any

_HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "item": (
        "item id",
        "item #",
        "item no",
        "item number",
        "item_no",
        "item",
        "line",
        "line #",
        "line no",
        "line number",
        "clin",
        "lineno",
    ),
    # Prefer specific product description headers over subcategory / group labels
    "description": (
        "item description",
        "product description",
        "material description",
        "commodity description",
        "description",
        "desc",
        "product",
        "commodity",
        "title",
    ),
    "manufacturer": ("manufacturer", "mfr", "mfg", "make", "brand", "vendor", "oem"),
    "model": ("model", "model #", "model no", "model number", "model_no", "model #"),
    "part_number": ("part number", "part no", "part #", "p/n", "pn", "mpn", "part_number"),
    "catalog_number": (
        "vendor catalog number",
        "catalog #",
        "catalog no",
        "catalog number",
        "cat #",
        "cat no",
        "cat.",
        "catalog",
    ),
    "sku": ("sku", "item sku", "stock #", "stock number"),
    "nsn": ("nsn", "national stock number"),
    "quantity": ("qty", "quantity", "qnty", "qty ordered", "est qty", "estimated qty"),
    "uom": ("unit measure", "uom", "unit of measure", "unit of issue", "units", "unit"),
    "unit_price": ("unit cost", "unit price", "price", "bid price", "unit $"),
    "extended_price": ("extended", "extended price", "ext price", "total", "amount", "line total"),
    "pack": ("pack", "pack size", "case pack", "units/case"),
}


def _norm_header(h: Any) -> str:
    s = re.sub(r"\s+", " ", str(h or "").strip().lower())
    s = s.replace("_", " ")
    return s


def map_headers(headers: list[Any]) -> dict[str, int]:
    """Map canonical field → column index.

    Prefers exact / longer alias matches so 'Item Description' wins over
    'Subcategory Description' (both contain the word description).
    """
    mapped: dict[str, int] = {}
    norms = [_norm_header(h) for h in headers]
    for field, aliases in _HEADER_ALIASES.items():
        best_i: int | None = None
        best_score = -1
        for i, n in enumerate(norms):
            if not n:
                continue
            for a in aliases:
                if n == a:
                    score = 200 + len(a)
                elif n.startswith(a + " ") or n.endswith(" " + a):
                    score = 120 + len(a)
                elif a in n:
                    score = 80 + len(a) - (len(n) - len(a))
                else:
                    continue
                if field == "description":
                    if n.startswith("item ") or n == "item description":
                        score += 40
                    if "subcategor" in n or n.startswith("group"):
                        score -= 50
                if field == "item" and "subcategor" in n:
                    score -= 40
                if score > best_score:
                    best_score = score
                    best_i = i
        if best_i is not None:
            mapped[field] = best_i
    return mapped


def _cell(row: list[Any], idx: int | None) -> Any:
    if idx is None or idx < 0 or idx >= len(row):
        return None
    v = row[idx]
    if v is None:
        return None
    if isinstance(v, str) and not v.strip():
        return None
    return v


def rows_from_matrix(
    matrix: list[list[Any]],
    *,
    source_path: str,
    sheet: str | None = None,
) -> list[dict[str, Any]]:
    """Detect header row and emit structured dict rows."""
    if not matrix:
        return []
    best_i = 0
    best_score = -1
    for i, row in enumerate(matrix[:40]):
        mapped = map_headers(row)
        score = len(mapped)
        if "description" in mapped:
            score += 2
        if "quantity" in mapped or "part_number" in mapped or "catalog_number" in mapped:
            score += 1
        if score > best_score:
            best_score = score
            best_i = i
    if best_score < 2:
        return []
    headers = matrix[best_i]
    mapped = map_headers(headers)
    # Track manufacturer header inheritance from section rows
    inherited_mfr: str | None = None
    out: list[dict[str, Any]] = []
    for ridx, row in enumerate(matrix[best_i + 1 :], start=best_i + 2):
        if not any(c not in (None, "") for c in row):
            continue
        # Section header: single non-empty cell looking like manufacturer/section
        nonempty = [c for c in row if c not in (None, "")]
        if len(nonempty) == 1 and isinstance(nonempty[0], str):
            text = nonempty[0].strip()
            if 2 <= len(text) <= 80 and not re.search(r"\b(qty|quantity|total|page)\b", text, re.I):
                inherited_mfr = text
                continue
        desc = _cell(row, mapped.get("description"))
        item = _cell(row, mapped.get("item"))
        pn = _cell(row, mapped.get("part_number"))
        cat = _cell(row, mapped.get("catalog_number"))
        model = _cell(row, mapped.get("model"))
        mfr = _cell(row, mapped.get("manufacturer"))
        qty = _cell(row, mapped.get("quantity"))
        # Skip spreadsheet type-code rows (e.g. C11 / C40 / N11.2 under BidNet forms)
        blob = " ".join(str(x) for x in (item, desc, qty, _cell(row, mapped.get("uom"))) if x is not None)
        if re.fullmatch(r"[A-Z]\d{1,4}([./]\d+)?", str(desc or "").strip()) and re.fullmatch(
            r"[A-Z]\d{1,4}", str(item or "").strip()
        ):
            continue
        if re.search(r"\bN\d+\.\d+\b", str(qty or "")) and re.fullmatch(
            r"[A-Z]\d{1,4}", str(item or "").strip() or "x"
        ):
            continue
        if desc is None and pn is None and cat is None and model is None:
            # Sometimes description is first free-text col
            continue
        if desc is None and (pn or cat or model):
            desc = str(pn or cat or model)
        row_d = {
            "item": item,
            "description": desc,
            "manufacturer": mfr or inherited_mfr,
            "model": model,
            "part_number": pn,
            "catalog_number": cat,
            "sku": _cell(row, mapped.get("sku")),
            "nsn": _cell(row, mapped.get("nsn")),
            "quantity": qty,
            "uom": _cell(row, mapped.get("uom")),
            "unit_price": _cell(row, mapped.get("unit_price")),
            "extended_price": _cell(row, mapped.get("extended_price")),
            "pack": _cell(row, mapped.get("pack")),
            "inherited_manufacturer": bool(inherited_mfr and not mfr),
            "_source_path": source_path,
            "_sheet": sheet,
            "_row": ridx,
            "_header_map": mapped,
        }
        out.append(row_d)
    return out


def extract_xlsx(path: Path) -> list[dict[str, Any]]:
    try:
        import openpyxl
    except Exception:
        return []
    rows: list[dict[str, Any]] = []
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return []
    try:
        for sn in wb.sheetnames:
            ws = wb[sn]
            matrix: list[list[Any]] = []
            for row in ws.iter_rows(values_only=True):
                matrix.append(list(row))
                if len(matrix) > 5000:
                    break
            rows.extend(rows_from_matrix(matrix, source_path=str(path), sheet=sn))
    finally:
        try:
            wb.close()
        except Exception:
            pass
    return rows


def extract_xls(path: Path) -> list[dict[str, Any]]:
    try:
        import xlrd
    except Exception:
        return []
    try:
        book = xlrd.open_workbook(str(path))
    except Exception:
        return []
    rows: list[dict[str, Any]] = []
    for si in range(book.nsheets):
        sh = book.sheet_by_index(si)
        matrix = [sh.row_values(r) for r in range(min(sh.nrows, 5000))]
        rows.extend(rows_from_matrix(matrix, source_path=str(path), sheet=sh.name))
    return rows


def extract_csv_file(path: Path) -> list[dict[str, Any]]:
    try:
        text = path.read_text(encoding="utf-8", errors="replace")
    except Exception:
        try:
            text = path.read_text(encoding="latin-1", errors="replace")
        except Exception:
            return []
    # Only treat as CSV if it has a plausible header line
    first = (text.splitlines() or [""])[0]
    if first.count(",") < 2 and first.count("\t") < 2:
        return []
    dialect = csv.excel
    if first.count("\t") > first.count(","):
        dialect = csv.excel_tab
    reader = csv.reader(text.splitlines(), dialect)
    matrix = [list(r) for r in reader]
    return rows_from_matrix(matrix, source_path=str(path), sheet="csv")


def extract_spreadsheet(path: Path) -> list[dict[str, Any]]:
    ext = path.suffix.lower()
    if ext == ".xlsx":
        return extract_xlsx(path)
    if ext == ".xls":
        return extract_xls(path)
    if ext == ".csv":
        return extract_csv_file(path)
    return []
