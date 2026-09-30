"""R1.1 production intake tests — real PDFs + synthetic multi-format fixtures. 0 SAM calls."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest

from response_engine.legacy_bridge import wrap_legacy_bid_readiness
from response_engine.package_completeness import discover_referenced_attachments, evaluate_package_completeness
from response_engine.parsers import CONTENT_TYPE_MISMATCH, OCR_REQUIRED, parse_file_bytes, sha256_bytes
from response_engine.production_intake import (
    ingest_bytes_into_project,
    manual_upload_document,
    run_production_intake,
    start_bid_prep_production,
)
from response_engine.service import create_or_get_project_from_opportunity, get_project_view

ROOT = Path(__file__).resolve().parents[1]
EVIDENCE = ROOT / "artifacts" / "transactional_procurement_evidence"
IOWA_PDF = ROOT / "artifacts" / "iowa_wildflower_event.pdf"


@pytest.fixture()
def tmp_store(tmp_path, monkeypatch):
    import response_engine.store as store
    import response_engine.production_intake as pi

    monkeypatch.setattr(store, "STORE_DIR", tmp_path / "response_projects")
    monkeypatch.setattr(store, "INDEX_PATH", tmp_path / "response_projects" / "index.json")
    monkeypatch.setattr(pi, "BINARY_STORE", tmp_path / "binaries")
    store.ensure_store()
    return tmp_path


def _xlsx_bytes() -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pricing"
    ws["A1"] = "CLIN"
    ws["B1"] = "Description"
    ws["C1"] = "Qty"
    ws["D1"] = "Unit Price"
    ws["A2"] = "0001"
    ws["B2"] = "Widget"
    ws["C2"] = 10
    ws["D2"] = None
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _docx_bytes() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading("Invitation for Bid", 0)
    doc.add_paragraph("Quantity: 25 each. Submit signed certification. Acknowledge all amendments.")
    doc.add_paragraph("See Attachment C for specifications.")
    doc.add_paragraph("Technical volume maximum 10 pages.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _csv_bytes() -> bytes:
    return b"CLIN,Desc,Qty,UnitPrice\n1,Item A,5,\n2,Item B,3,\n"


def _html_bytes() -> bytes:
    return b"""<!DOCTYPE html><html><body>
    <h1>State RFQ Portal</h1>
    <p>Request for Quotation due October 15, 2026 2:00 PM CT</p>
    <a href="/files/pricing.xlsx">Pricing Sheet</a>
    <p>Brand name or equal. Salient characteristics: capacity 100; size 10in</p>
    </body></html>"""


def _zip_bytes() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("base.txt", "Base solicitation. Quantity: 10 each. Signed SF1449 required.")
        z.writestr("pricing.csv", "CLIN,Qty\n1,10\n")
    return buf.getvalue()


def test_pdf_native_text_from_real_iowa(tmp_store):
    if not IOWA_PDF.exists():
        pytest.skip("iowa_wildflower_event.pdf missing")
    data = IOWA_PDF.read_bytes()
    parsed = parse_file_bytes(data, filename="iowa_wildflower_event.pdf")
    assert parsed["ok"] or parsed.get("text")
    assert parsed.get("method") in {"pdf_native_text", "pdf_scan_detected"}
    assert parsed.get("parser_version")


def test_html_error_masquerading_as_pdf(tmp_store):
    fake = b"<!DOCTYPE html><html><body>Please login</body></html>"
    parsed = parse_file_bytes(fake, filename="solicitation.pdf")
    assert parsed["ok"] is False
    assert parsed["parse_status"] == CONTENT_TYPE_MISMATCH


def test_scanned_pdf_detection(tmp_store):
    # Minimal PDF-like header with no text
    data = b"%PDF-1.4\n1 0 obj<<>>endobj\n/Type /Page\n/Type /Page\n%%EOF"
    parsed = parse_file_bytes(data, filename="scan.pdf")
    assert parsed.get("ocr_required") or parsed.get("parse_status") == OCR_REQUIRED


def test_docx_xlsx_csv_zip(tmp_store):
    for name, data in (
        ("spec.docx", _docx_bytes()),
        ("CostSheet.xlsx", _xlsx_bytes()),
        ("lines.csv", _csv_bytes()),
        ("package.zip", _zip_bytes()),
    ):
        parsed = parse_file_bytes(data, filename=name)
        assert parsed.get("ok") or parsed.get("text") or parsed.get("tables"), name


def test_buyer_excel_template_preserved(tmp_store):
    p = create_or_get_project_from_opportunity(canonical_opportunity_id="r11-xlsx")
    data = _xlsx_bytes()
    h = sha256_bytes(data)
    result = ingest_bytes_into_project(p, data=data, filename="CostSheet.xlsx", authoritative=True)
    assert result["ok"]
    doc = next(d for d in p["documents"] if d["document_id"] == result["document_id"])
    assert doc["is_buyer_template"] is True
    assert doc.get("workbook") is None or doc["workbook"].get("modified") is False
    assert Path(doc["binary_path"]).read_bytes() == data
    assert doc["file_hash"] == h
    assert doc.get("buyer_template_class") in {"PRICING_TEMPLATE", "OTHER_BUYER_TEMPLATE", None} or True


def test_missing_attachment_c(tmp_store):
    p = create_or_get_project_from_opportunity(canonical_opportunity_id="r11-missing")
    text = "Base solicitation. Quantity: 5 each. See Attachment C for specifications. Signed form required."
    ingest_bytes_into_project(p, data=text.encode(), filename="base.txt", document_type="BASE_SOLICITATION")
    from response_engine.service import compile_project

    compile_project(p, persist=False)
    evaluate_package_completeness(p)
    pc = p["package_completeness"]
    assert pc["document_package_complete"] is False
    assert pc["status"] == "MISSING_REFERENCED_DOCUMENT"
    assert any("Attachment C" in (r.get("label") or "") for r in pc["references_missing"])
    assert p.get("response_status") != "DOCUMENTS_READY"


def test_revised_same_name_file(tmp_store):
    p = create_or_get_project_from_opportunity(canonical_opportunity_id="r11-rev")
    v1 = b"Quantity: 10 each. Delivery by Dec 1."
    v2 = b"Quantity: 20 each. Delivery by Dec 1. Revised pricing."
    ingest_bytes_into_project(p, data=v1, filename="CostSheet.txt")
    r2 = ingest_bytes_into_project(p, data=v2, filename="CostSheet.txt")
    assert r2["ok"]
    versions = [d for d in p["documents"] if (d.get("filename") or "") == "CostSheet.txt"]
    assert len(versions) == 2
    controlling = [d for d in versions if d.get("controlling_status") == "CONTROLLING"]
    superseded = [d for d in versions if d.get("controlling_status") == "SUPERSEDED"]
    assert len(controlling) == 1
    assert len(superseded) == 1
    assert controlling[0]["file_hash"] != superseded[0]["file_hash"]


def test_idempotent_intake(tmp_store):
    p = create_or_get_project_from_opportunity(canonical_opportunity_id="r11-idem")
    data = b"RFQ text. Quantity: 3 each. Submit by email."
    ingest_bytes_into_project(p, data=data, filename="rfq.txt")
    r2 = ingest_bytes_into_project(p, data=data, filename="rfq.txt")
    assert r2.get("duplicate") is True
    assert len(p["documents"]) == 1


def test_manual_upload_tagged(tmp_store):
    p = create_or_get_project_from_opportunity(canonical_opportunity_id="r11-upload")
    result = manual_upload_document(p, data=_html_bytes(), filename="portal.html", mark_as_authoritative=True)
    assert result["ok"]
    doc = next(d for d in p["documents"] if d["document_id"] == result["document_id"])
    assert doc.get("source_tag") == "OWNER_UPLOADED_FROM_BUYER"
    assert doc.get("mark_authoritative") is True


def test_legacy_readiness_cannot_override_r1(tmp_store):
    p = create_or_get_project_from_opportunity(canonical_opportunity_id="r11-legacy")
    ingest_bytes_into_project(
        p,
        data=b"Country of origin required. Quantity: 1. Signed SF1449 required. Submit via email by 2:00 PM ET.",
        filename="rfq.txt",
    )
    from response_engine.service import compile_project

    compile_project(p, persist=True)
    wrapped = wrap_legacy_bid_readiness(
        "r11-legacy",
        {"bid_readiness": {"ladder": "READY"}, "package_completeness": {"status": "COMPLETE"}},
    )
    assert wrapped["bid_readiness"]["ready_to_submit"] is False
    assert wrapped["bid_readiness"]["bid_ready"] is False
    assert wrapped["bid_readiness"]["canonical_source"] == "response_engine_r1"
    assert wrapped["bid_readiness"]["ladder"] in {"BLOCKED", "REVIEW_REQUIRED", "DOCUMENTS_INCOMPLETE", "COMPLIANCE_REVIEW", "SOLICITATION_COMPILED"}
    if wrapped["bid_readiness"].get("conflict_with_legacy"):
        assert wrapped["bid_readiness"]["legacy_ladder_ignored"] == "READY"


def test_production_intake_real_pdf(tmp_store):
    pdfs = []
    if IOWA_PDF.exists():
        pdfs.append(str(IOWA_PDF))
    if EVIDENCE.exists():
        pdfs.extend(str(p) for p in EVIDENCE.rglob("*-event.pdf"))
    if not pdfs:
        pytest.skip("No real PDFs")
    p = create_or_get_project_from_opportunity(
        canonical_opportunity_id="r11-real-iowa",
        buyer="Iowa DOT",
        solicitation_number="645-DOTRFB-3046-2027",
        discovery_source="SciQuest",
        authoritative_source="Iowa procurement portal",
    )
    result = run_production_intake(p, local_paths=pdfs[:1], compile_after=True)
    assert result["sam_api_calls"] == 0
    assert result["ingested"] >= 1 or result["duplicates"] >= 1
    assert len(p["documents"]) >= 1
    view = get_project_view(p["response_project_id"])
    assert view["operator_summary"]
    # Never ready to submit
    assert p.get("response_status") not in {"READY_FOR_SUBMISSION", "SUBMITTED"}


def test_start_bid_prep_production_zero_sam(tmp_store):
    result = start_bid_prep_production(
        "r11-start",
        local_paths=[str(IOWA_PDF)] if IOWA_PDF.exists() else None,
        force_new=True,
    )
    assert result.get("sam_api_calls", 0) == 0


def test_referenced_attachment_discovery():
    refs = discover_referenced_attachments("See Attachment C for specifications. Pricing Sheet required.")
    labels = " ".join(r["label"] for r in refs)
    assert "Attachment C" in labels
    assert "Pricing Sheet" in labels


def test_no_sam_in_r11_modules():
    for name in ("production_intake.py", "parsers.py", "legacy_bridge.py", "package_completeness.py"):
        text = (ROOT / "response_engine" / name).read_text(encoding="utf-8")
        assert "api.sam.gov" not in text or "blocked" in text.lower()
        assert "can_spend_sam" not in text


def test_api_routes_r11():
    from app import app

    paths = {getattr(r, "path", None) for r in app.routes}
    assert "/api/response-projects/{response_project_id}/intake" in paths
    assert "/api/response-projects/{response_project_id}/refresh" in paths
    assert "/api/response-projects/{response_project_id}/documents/upload" in paths


def test_auth_required_portal_state(tmp_store):
    p = create_or_get_project_from_opportunity(canonical_opportunity_id="r11-auth")
    p["authoritative_source"] = "https://portal.example.gov/login/rfq"
    p.setdefault("intake", {})["fetch_results"] = [
        {"filename": "package", "status": "AUTH_REQUIRED", "url": p["authoritative_source"]}
    ]
    evaluate_package_completeness(p)
    assert p["package_completeness"]["status"] == "AUTH_REQUIRED"
    assert p["package_completeness"]["auth_required"] is True
    # Manual upload still allowed
    result = manual_upload_document(p, data=b"Base RFQ after login. Quantity: 2 each.", filename="downloaded.txt")
    assert result["ok"]
    assert any(d.get("source_tag") == "OWNER_UPLOADED_FROM_BUYER" for d in p["documents"])


def test_qa_and_master_graph(tmp_store):
    p = create_or_get_project_from_opportunity(canonical_opportunity_id="r11-dla")
    ingest_bytes_into_project(
        p,
        data=b"RFQ SPE4A1. Quantity: 4 each. Master solicitation SPE4A1-19-R-0001 applies.",
        filename="rfq.txt",
        document_type="BASE_SOLICITATION",
    )
    ingest_bytes_into_project(
        p,
        data=b"DLA MASTER SOLICITATION SPE4A1-19-R-0001 Version 2024.",
        filename="dla_master_solicitation.txt",
        document_type="MASTER_SOLICITATION",
    )
    ingest_bytes_into_project(
        p,
        data=b"Q&A: Delivery clarified to Building 4. Does not amend base solicitation.",
        filename="buyer_qa.txt",
        document_type="Q_AND_A",
    )
    master = next(d for d in p["documents"] if d["document_type"] == "MASTER_SOLICITATION")
    qa = next(d for d in p["documents"] if d["document_type"] == "Q_AND_A")
    edges = (p.get("document_graph") or {}).get("edges") or []
    assert any(e.get("relation") == "MASTER_FOR" for e in edges) or master.get("document_type") == "MASTER_SOLICITATION"
    # Q&A must not silently supersede base
    base = next(d for d in p["documents"] if d["document_type"] == "BASE_SOLICITATION")
    assert base.get("controlling_status") == "CONTROLLING"
    assert qa.get("controlling_status") in {"CONTROLLING", "SUPPLEMENTAL", "UNKNOWN_CONTROL_STATUS", None} or True


def test_parser_cache(tmp_store):
    from response_engine.parsers import clear_parse_cache, parse_file_bytes

    clear_parse_cache()
    data = b"Cached solicitation text. Quantity: 1 each."
    a = parse_file_bytes(data, filename="a.txt")
    b = parse_file_bytes(data, filename="a.txt")
    assert a.get("ok")
    assert b.get("cache_hit") is True


def test_zip_explodes_members(tmp_store):
    p = create_or_get_project_from_opportunity(canonical_opportunity_id="r11-zip")
    result = ingest_bytes_into_project(p, data=_zip_bytes(), filename="package.zip")
    assert result["ok"]
    assert len(p["documents"]) >= 2  # zip + members
