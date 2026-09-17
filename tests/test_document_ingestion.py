"""Tests for authorized document ingestion, reprocess, blockers, bid targets."""

from __future__ import annotations

import io
import json
import zipfile
from pathlib import Path

import pytest

from bid_price_targets import compute_bid_price_targets
from document_ingestion import (
    ingest_solicitation_document,
    load_registry,
    normalize_acquisition_method,
)
from document_ingestion_constants import (
    ACQ_AUTHORIZED_OPERATOR_DOWNLOAD,
    ACQ_MANUAL_UPLOAD,
    BID_BREAK_EVEN,
    BID_MIN_10K,
    BLOCKER_AUTH,
    BLOCKER_FUNDING,
    BLOCKER_OPERATOR_DOC,
    BLOCKER_SUPPLIER_QUOTE,
    FIT_QUOTE,
)
from procurement_blockers import classify_procurement_blockers
from procurement_reprocess import diff_snapshots, reprocess_solicitation
from specification_extraction import extract_blade_or_product_specifications
from supplier_fit import evaluate_supplier_fit
from table_extractors import safe_inspect_zip


BLADE_SPEC_TEXT = """
Snow Plow Blade Specification
Iowa DOT

Blade length: 3 ft and 4 ft sections
Blade width: 6 inches
Blade thickness: 0.875 inches
Tungsten-carbide insert configuration: carbide insert rubber encased
Rubber encased cutting edge required
Hole pattern: square holes per cover strap drawing
Mounting requirements: bolt to moldboard with cover straps and back supports
Material: high-carbon steel with tungsten-carbide insert
Hardness: Rockwell C 60 minimum on carbide
Performance: suitable for snow/ice removal on state highways
Cover strap requirements: 3 FT and 4 FT cover straps with square holes
Back support requirements: 3 FT and 4 FT back supports
Approved brands: any comparable brand or equal — not intended to be restrictive
Mill certifications required with shipment
Samples may be required within five business days
Testing: visual inspection and mill cert review
Packaging: banded on 40x48 pallets not exceeding 60 inches height
Deliveries shall be F.O.B Destination
Warranty: manufacturer standard warranty
Country of origin must be disclosed
"""


@pytest.fixture()
def tmp_sol(tmp_path, monkeypatch):
    sol = "645-DOTRFB-2975-2027"
    # Redirect inbox/registry into tmp
    import document_ingestion as di
    import procurement_reprocess as pr

    monkeypatch.setattr(di, "INBOX_ROOT", tmp_path / "inbox")
    monkeypatch.setattr(di, "REGISTRY_ROOT", tmp_path / "registry")
    monkeypatch.setattr(di, "ARTIFACTS", tmp_path / "artifacts")
    monkeypatch.setattr(di, "CHANGE_HISTORY_PATH", tmp_path / "artifacts" / "document_change_history.json")
    monkeypatch.setattr(pr, "ARTIFACTS", tmp_path / "artifacts")
    monkeypatch.setattr(pr, "PACKETS_DIR", tmp_path / "packets")
    monkeypatch.setattr(pr, "EVIDENCE_DIR", tmp_path / "evidence")
    (tmp_path / "packets").mkdir()
    (tmp_path / "artifacts").mkdir()
    (tmp_path / "evidence" / sol).mkdir(parents=True)
    return sol, tmp_path


def _write_blade_packet(packets_dir: Path, sol: str) -> dict:
    packet = {
        "solicitation_number": sol,
        "title": "Tungsten-Carbide Blades for Snow/Ice Removal",
        "agency": "State of Iowa Department of Administrative Services",
        "deal_type": "ONE_TIME_PRODUCT_PURCHASE",
        "documents": [
            {
                "document_title": f"{sol}-event.pdf",
                "access_status": "PUBLIC_FETCHED",
                "document_class": "SOLICITATION",
                "hash": "aaa",
            },
            {
                "document_title": "Snow Plow Blade Spec 8Sept20263.pdf",
                "access_status": "LOGIN_REQUIRED",
                "document_class": "SPECIFICATION",
            },
        ],
        "requirement": {
            "line_items": [
                {
                    "line_number": "P1",
                    "CLIN_or_item_number": "002367300",
                    "description": "Tungsten-Carbide BLADE 3 FT. SECTION (rubber encased)",
                    "quantity": 300.0,
                    "unit_of_measure": "EA",
                },
                {
                    "line_number": "P2",
                    "description": "Tungsten-Carbide BLADE 4 FT. SECTION (rubber encased)",
                    "quantity": 600.0,
                    "unit_of_measure": "EA",
                    "CLIN_or_item_number": "002367400",
                },
            ],
            "terms": {
                "delivery_location": {"value": "931 S. 4th Street, Ames, IA 50010"},
                "bid_deadline": {"value": "10/7/2026, 1:00 PM CDT"},
                "brand_name_or_equal": {"value": True},
                "listed_buyer_attachments": {"value": ["Snow Plow Blade Spec 8Sept20263.pdf"]},
            },
            "product_identification": {
                "product_id_state": "BRAND_OR_EQUAL_IDENTIFIED",
                "sourcing_description": "Tungsten-Carbide BLADE 3 FT / 4 FT",
            },
            "completeness": {
                "document_completeness": "COMPLETE_ENOUGH_FOR_PRODUCT_ID",
                "critical_spec_missing": True,
                "missing": ["specification_attachment:Snow Plow Blade Spec 8Sept20263.pdf"],
            },
        },
        "suppliers": [],
        "economics_block": {
            "supplier_cost": None,
            "freight_status": "UNKNOWN",
            "expected_revenue": None,
            "profit_state": "PROFIT_UNKNOWN",
            "working_capital_status": "UNKNOWN",
            "funding_warranted": False,
        },
        "funding": {"status": "NOT_YET_NEEDED"},
        "compliance": {"entity_eligibility": "UNKNOWN", "set_aside": "UNKNOWN"},
        "government_value_evidence": {"notes": "none"},
        "bidder_priced": True,
    }
    (packets_dir / f"{sol}.json").write_text(json.dumps(packet), encoding="utf-8")
    return packet


def test_normalize_acquisition_and_manual_not_weak():
    assert normalize_acquisition_method("authorized_operator_download") == ACQ_AUTHORIZED_OPERATOR_DOWNLOAD
    assert normalize_acquisition_method("manual_upload") == ACQ_MANUAL_UPLOAD


def test_manual_pdf_ingestion_hash_and_provenance(tmp_sol):
    sol, tmp = tmp_sol
    pdf = tmp / "Snow Plow Blade Spec 8Sept20263.pdf"
    # Minimal PDF header enough for storage; text via sidecar path uses extract — write as txt ingested with .pdf name won't extract text well
    # Use .txt classified as SPECIFICATION for extraction path, and also a real-ish pdf bytes + companion
    pdf.write_bytes(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF\n")
    # Also write a text file that will be the real content ingest for classification tests
    txt = tmp / "Snow_Plow_Blade_Spec.txt"
    txt.write_text(BLADE_SPEC_TEXT, encoding="utf-8")

    r1 = ingest_solicitation_document(
        solicitation_number=sol,
        file_path=txt,
        source_system="sciquest",
        acquisition_method="authorized_operator_download",
        document_title="Snow Plow Blade Spec 8Sept20263.pdf",
        document_type="SPECIFICATION",
    )
    assert r1["ok"] and r1["ingested"]
    doc = r1["document"]
    assert doc["hash"]
    assert doc["acquisition_method"] == ACQ_AUTHORIZED_OPERATOR_DOWNLOAD
    assert doc["appears_authoritative"] is True
    assert doc["solicitation_number"] == sol
    assert doc["document_class"] == "SPECIFICATION"

    # Duplicate detection
    r2 = ingest_solicitation_document(
        solicitation_number=sol,
        file_path=txt,
        acquisition_method="manual_upload",
        document_title="Snow Plow Blade Spec 8Sept20263.pdf",
        document_type="SPECIFICATION",
    )
    assert r2["duplicate"] is True
    assert r2["ingested"] is False

    reg = load_registry(sol)
    assert len(reg["documents"]) == 1


def test_docx_xlsx_csv_zip_ingestion(tmp_sol):
    sol, tmp = tmp_sol
    from docx import Document
    import openpyxl

    # DOCX
    doc = Document()
    doc.add_paragraph("Specification hole pattern square holes")
    docx_path = tmp / "spec.docx"
    doc.save(docx_path)
    r = ingest_solicitation_document(
        solicitation_number=sol,
        file_path=docx_path,
        acquisition_method="manual_upload",
        document_type="SPECIFICATION",
    )
    assert r["ok"] and r["ingested"]

    # XLSX
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Line", "Description", "Qty", "UOM", "Unit Price"])
    ws.append([1, "Blade", 10, "EA", ""])
    xlsx = tmp / "bid.xlsx"
    wb.save(xlsx)
    r = ingest_solicitation_document(
        solicitation_number=sol, file_path=xlsx, acquisition_method="manual_upload", document_type="BID_SCHEDULE"
    )
    assert r["ok"] and r["ingested"]

    # CSV
    csvp = tmp / "lines.csv"
    csvp.write_text("line,description,quantity,uom,unit_price\n1,A,2,EA,\n", encoding="utf-8")
    r = ingest_solicitation_document(
        solicitation_number=sol, file_path=csvp, acquisition_method="manual_upload", document_type="LINE_ITEM_SCHEDULE"
    )
    assert r["ok"] and r["ingested"]

    # ZIP safe
    zpath = tmp / "pkg.zip"
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("inner.txt", "carbide blade width: 6 inches")
    zpath.write_bytes(buf.getvalue())
    r = ingest_solicitation_document(
        solicitation_number=sol, file_path=zpath, acquisition_method="authorized_portal_export"
    )
    assert r["ok"] and r["ingested"]


def test_zip_traversal_rejection():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("../evil.txt", "x")
        zf.writestr("ok.txt", "y")
    result = safe_inspect_zip(buf.getvalue())
    assert any(r.get("reason") == "path_traversal_or_absolute" for r in result["rejected"])


def test_version_preservation_and_amendment_precedence(tmp_sol):
    sol, tmp = tmp_sol
    f1 = tmp / "spec_v1.txt"
    f1.write_text("version one thickness: 0.5 inches", encoding="utf-8")
    f2 = tmp / "spec_v2.txt"
    f2.write_text("version two thickness: 0.875 inches amendment", encoding="utf-8")
    ingest_solicitation_document(
        solicitation_number=sol,
        file_path=f1,
        document_title="Snow Spec.pdf",
        document_type="SPECIFICATION",
        acquisition_method="manual_upload",
    )
    ingest_solicitation_document(
        solicitation_number=sol,
        file_path=f2,
        document_title="Snow Spec.pdf",
        document_type="AMENDMENT",
        acquisition_method="manual_upload",
        amendment_sequence=2,
    )
    reg = load_registry(sol)
    assert len(reg["documents"]) == 2
    assert sum(1 for d in reg["documents"] if d.get("superseded")) == 1
    assert any(not d.get("superseded") and d.get("document_class") == "AMENDMENT" for d in reg["documents"])


def test_blade_specification_extraction_and_unknown_preservation():
    spec = extract_blade_or_product_specifications(BLADE_SPEC_TEXT, source_document="spec.pdf")
    fields = spec["specifications"]
    assert fields["blade_width"]["value"]
    assert fields["carbide_configuration"]["value"]
    assert fields["mill_certifications"]["value"] is True
    assert fields["brand_or_equal_language"]["value"] is True
    # Unknown invented fields stay unknown when absent
    empty = extract_blade_or_product_specifications("No technical content here.", source_document="x")
    assert empty["specifications"]["hardness"]["value"] is None
    assert empty["specifications"]["hardness"]["confidence"] == "UNKNOWN"


def test_blocker_auth_and_funding_not_premature():
    packet = {
        "documents": [
            {
                "document_title": "Snow Plow Blade Spec 8Sept20263.pdf",
                "access_status": "LOGIN_REQUIRED",
                "document_class": "SPECIFICATION",
            }
        ],
        "ingested_documents": [],
        "requirement": {
            "product_identification": {"product_id_state": "BRAND_OR_EQUAL_IDENTIFIED"},
            "completeness": {"critical_spec_missing": True},
        },
        "economics_block": {
            "supplier_cost": None,
            "freight_status": "UNKNOWN",
            "expected_revenue": None,
            "funding_warranted": False,
        },
        "funding": {"status": "NOT_YET_NEEDED"},
        "compliance": {"entity_eligibility": "UNKNOWN"},
        "suppliers": [{"supplier_name": "X"}],
        "bidder_priced": True,
        "government_value_evidence": {},
    }
    bm = classify_procurement_blockers(packet)
    assert BLOCKER_AUTH in bm["codes"]
    assert BLOCKER_OPERATOR_DOC in bm["codes"]
    assert BLOCKER_SUPPLIER_QUOTE in bm["codes"]
    assert BLOCKER_FUNDING not in bm["codes"]
    assert bm["funding_premature"] is True
    assert bm["funding_is_immediate_blocker"] is False
    assert bm["primary_blocker"] in {BLOCKER_AUTH, BLOCKER_OPERATOR_DOC}


def test_automatic_reprocess_before_after(tmp_sol):
    sol, tmp = tmp_sol
    _write_blade_packet(tmp / "packets", sol)
    # Seed evidence text with line items so reprocess keeps BOM
    (tmp / "evidence" / sol / "event.txt").write_text(
        """
Product Line Items
P1
002367300 - Tungsten-Carbide BLADE 3 FT. SECTION
(rubber encased)
300
EA - Each
P2
002367400 - Tungsten-Carbide BLADE 4 FT. SECTION
(rubber encased)
600
EA - Each
Service Line Items
There are no Items added to this event.
Close
10/7/2026, 1:00 PM CDT
Deliveries shall be F.O.B Destination. Address: 931 S. 4th Street
Ames, IA 50010
""",
        encoding="utf-8",
    )
    before_packet = json.loads((tmp / "packets" / f"{sol}.json").read_text(encoding="utf-8"))
    spec = tmp / "Snow Plow Blade Spec 8Sept20263.txt"
    spec.write_text(BLADE_SPEC_TEXT, encoding="utf-8")
    ingest_solicitation_document(
        solicitation_number=sol,
        file_path=spec,
        document_title="Snow Plow Blade Spec 8Sept20263.pdf",
        document_type="SPECIFICATION",
        acquisition_method="authorized_operator_download",
    )
    result = reprocess_solicitation(sol, write_artifacts=True)
    assert result["ok"]
    assert result["packet"]["prior_evidence_preserved"] is True
    # Spec known fields should appear / blockers should drop AUTH if resolved
    codes = result["blocker_model"]["codes"]
    assert BLOCKER_AUTH not in codes
    assert BLOCKER_OPERATOR_DOC not in codes
    assert BLOCKER_SUPPLIER_QUOTE in codes
    assert any(c["field"] for c in result["changes"]) or result["packet"].get("technical_specifications")
    tech = result["packet"]["technical_specifications"]["specifications"]
    assert tech["blade_width"]["value"]
    assert result["supplier_outreach"] == 0
    assert result["lender_outreach"] == 0
    assert result["bid_submissions"] == 0
    # before snapshot had critical missing; after should improve completeness or specs
    assert before_packet["requirement"]["completeness"]["critical_spec_missing"] is True


def test_supplier_fit_quote_required():
    suppliers = [
        {"supplier_name": "Winter Equipment Company", "supplier_type": "SPECIALTY_SUPPLIER", "notes": "snowplow blades"},
        {"supplier_name": "Kennametal Inc.", "supplier_type": "MANUFACTURER", "notes": "carbide"},
    ]
    spec = extract_blade_or_product_specifications(BLADE_SPEC_TEXT)
    fitted = evaluate_supplier_fit(
        suppliers=suppliers,
        specifications=spec,
        product_id={"sourcing_description": "tungsten-carbide blade"},
    )
    assert len(fitted) == 2
    assert all(f.get("quote_required") for f in fitted)
    assert all(f.get("outreach_performed") is False for f in fitted)
    assert any(f.get("fit_state") == FIT_QUOTE or f.get("fit_category") for f in fitted)


def test_bidder_priced_bid_thresholds_no_invented_win():
    # Without cost → UNKNOWN thresholds
    t0 = compute_bid_price_targets()
    assert t0["winning_price_invented"] is False
    assert all(x["is_winning_price_claim"] is False for x in t0["thresholds"])

    t1 = compute_bid_price_targets(supplier_cost=20000, freight=2000, target_margin_pct=20)
    kinds = {x["kind"]: x for x in t1["thresholds"]}
    assert kinds[BID_BREAK_EVEN]["value"] == 22000
    assert kinds[BID_MIN_10K]["value"] == 32000
    assert kinds[BID_MIN_10K]["is_winning_price_claim"] is False
    assert kinds["MINIMUM_BID_FOR_TARGET_MARGIN"]["value"] == 27500.0
    assert t1["bid_submissions"] == 0


def test_diff_snapshots_unknown_preservation():
    changes = diff_snapshots(
        {"product_id_state": "INSUFFICIENT_INFORMATION", "supplier_cost": None},
        {"product_id_state": "SPEC_DEFINED_COMMODITY", "supplier_cost": None},
    )
    assert any(c["field"] == "product_id_state" and c["after"] == "SPEC_DEFINED_COMMODITY" for c in changes)
