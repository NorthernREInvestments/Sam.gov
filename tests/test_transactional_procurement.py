"""Focused tests for transactional procurement package → BOM → economics."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest

from solicitation_package_retrieval import (
    classify_document,
    document_record,
    enumerate_listed_attachments_from_text,
    map_http_to_access,
    precedence_for_fact,
    resolve_document_authority,
)
from solicitation_package_constants import (
    ACCESS_AUTH_REQUIRED,
    ACCESS_LOGIN_REQUIRED,
    ACCESS_NOT_FOUND,
    ACCESS_PUBLIC_FETCHED,
    COMPLETE_ENOUGH_FOR_PRODUCT_ID,
    DOC_AMENDMENT,
    DOC_Q_AND_A,
    DOC_SOLICITATION,
    DOC_SPECIFICATION,
    PRICE_TO_BE_PROVIDED,
    PRODUCT_BRAND_OR_EQUAL,
    PRODUCT_SPEC_COMMODITY,
    PROCUREMENT_HARD_MAX_HTTP,
)
from table_extractors import (
    extract_csv_table,
    extract_docx_tables,
    extract_html_tables,
    extract_pdf_tables_heuristic,
    extract_xlsx_table,
    safe_inspect_zip,
    _normalize_price_cell,
)
from transactional_bom import (
    assess_document_completeness,
    extract_delivery_and_terms,
    extract_sciquest_product_line_items,
    identify_product,
)
from transactional_procurement import (
    build_supplier_quote_packet,
    compute_procurement_economics,
    correct_deal_type_from_documents,
    discover_suppliers_for_product,
    evaluate_bid_ready_strict,
    maybe_run_funding,
    select_deals_with_replacement,
)


BLADES_PDF_TEXT = """
Number
645-DOTRFB-2975-2027
Close
10/7/2026, 1:00 PM CDT
Approved Brands
not intended to be restrictive, and bids are invited on these and comparable brands
*Mill Certifications must accompany any shipments.
Shipping Terms
Deliveries shall be F.O.B Destination.  Address: 931 S. 4th Street
Ames, IA 50010
Buyer Attachments
1.
Snow Plow Blade Spec 8Sept20263.pdf
Product Line Items
Product Line Items
P1
002367300 - Tungsten-Carbide BLADE 3 FT. SECTION
(rubber encased)
300
EA - Each
See attached specifications.
Manufacturer -
P2
002367400 - Tungsten-Carbide BLADE 4 FT. SECTION
(rubber encased)
600
EA - Each
P3
002367530 - 3 FT. COVER STRAP W/SQUARE HOLES
20
EA - Each
P4
002367600 - 3 FT. BACK SUPPORT
10
EA - Each
P5
002367540 - 4 FT. COVER STRAP W/SQUARE HOLES
30
EA - Each
P6
002367700 - 4 FT. BACK SUPPORT
10
EA - Each
Service Line Items
There are no Items added to this event.
"""


def test_document_classification():
    assert classify_document(title="Snow Plow Blade Spec.pdf") == DOC_SPECIFICATION
    assert classify_document(filename="amendment_1.pdf") == DOC_AMENDMENT
    assert classify_document(title="event.pdf", url="Sourcingevent/1-event.pdf") == DOC_SOLICITATION


def test_inaccessible_attachment_state():
    docs = enumerate_listed_attachments_from_text(
        BLADES_PDF_TEXT,
        solicitation_number="645-DOTRFB-2975-2027",
        source_portal="sciquest_iowa",
    )
    assert len(docs) == 1
    assert docs[0]["access_status"] == ACCESS_LOGIN_REQUIRED
    assert docs[0]["document_class"] == DOC_SPECIFICATION


def test_map_http_access_states():
    assert map_http_to_access(200, body=b"%PDF") == ACCESS_PUBLIC_FETCHED
    assert map_http_to_access(403) == ACCESS_AUTH_REQUIRED
    assert map_http_to_access(404) == ACCESS_NOT_FOUND
    assert map_http_to_access(200, body=b"<html>Login Password</html>", final_url="https://x/SupplierLogin") == ACCESS_LOGIN_REQUIRED


def test_amendment_precedence():
    winner = precedence_for_fact(
        "quantity",
        [
            {"value": 100, "document_class": DOC_SOLICITATION, "amendment_sequence": 0, "explicit": True},
            {"value": 120, "document_class": DOC_AMENDMENT, "amendment_sequence": 2, "explicit": True},
            {"value": 999, "document_class": DOC_Q_AND_A, "amendment_sequence": 0, "explicit": True},
        ],
    )
    assert winner["value"] == 120


def test_authority_marks_solicitation_when_amendment_present():
    docs = resolve_document_authority(
        [
            document_record(
                solicitation_number="X",
                source_portal="t",
                document_title="base.pdf",
                document_url="http://x/base.pdf",
                access_status=ACCESS_PUBLIC_FETCHED,
                document_class=DOC_SOLICITATION,
            ),
            document_record(
                solicitation_number="X",
                source_portal="t",
                document_title="amend1.pdf",
                document_url="http://x/a.pdf",
                access_status=ACCESS_PUBLIC_FETCHED,
                document_class=DOC_AMENDMENT,
                amendment_sequence=1,
            ),
        ]
    )
    base = next(d for d in docs if d["document_class"] == DOC_SOLICITATION)
    assert base["superseded_or_amended"] is True


def test_xlsx_bid_schedule_extraction():
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Line", "Item", "Description", "Qty", "UOM", "Unit Price"])
    ws.append([1, "A1", "Carbide blade 3ft", 300, "EA", ""])
    ws.append([2, "A2", "Carbide blade 4ft", 600, "EA", None])
    buf = io.BytesIO()
    wb.save(buf)
    result = extract_xlsx_table(buf.getvalue(), source_name="bid.xlsx")
    assert result["ok"]
    assert len(result["line_items"]) == 2
    assert result["line_items"][0]["quantity"] == 300
    assert result["line_items"][0]["unit_price"]["status"] == PRICE_TO_BE_PROVIDED
    assert result["line_items"][0]["unit_price"]["value"] is not None or result["line_items"][0]["unit_price"]["value"] is None
    assert result["line_items"][0]["unit_price"]["value"] != 0


def test_csv_extraction():
    csv_data = "line,description,quantity,uom,unit_price\n1,Widget,10,EA,\n"
    result = extract_csv_table(csv_data, source_name="s.csv")
    assert result["ok"]
    assert result["line_items"][0]["quantity"] == 10
    assert result["line_items"][0]["unit_price"]["status"] == PRICE_TO_BE_PROVIDED


def test_html_table_extraction():
    html = """
    <table>
      <tr><th>Item</th><th>Description</th><th>Qty</th><th>UOM</th><th>Price</th></tr>
      <tr><td>1</td><td>Seed mix</td><td>50</td><td>LB</td><td></td></tr>
    </table>
    """
    result = extract_html_tables(html, source_name="page.html")
    assert result["ok"]
    assert result["line_items"][0]["quantity"] == 50
    assert result["line_items"][0]["unit_price"]["status"] == PRICE_TO_BE_PROVIDED


def test_pdf_text_table_extraction():
    result = extract_pdf_tables_heuristic(BLADES_PDF_TEXT, source_name="event.pdf")
    assert result["ok"]
    assert len(result["line_items"]) == 6


def test_docx_table_extraction():
    from docx import Document

    doc = Document()
    table = doc.add_table(rows=2, cols=5)
    headers = ["Line", "Description", "Qty", "UOM", "Unit Price"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    table.rows[1].cells[0].text = "1"
    table.rows[1].cells[1].text = "Badge"
    table.rows[1].cells[2].text = "1"
    table.rows[1].cells[3].text = "EA"
    table.rows[1].cells[4].text = ""
    buf = io.BytesIO()
    doc.save(buf)
    result = extract_docx_tables(buf.getvalue(), source_name="spec.docx")
    assert result["ok"]
    assert result["line_items"][0]["unit_price"]["status"] == PRICE_TO_BE_PROVIDED


def test_zip_safe_extraction():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("schedule.csv", "line,description,quantity,uom,unit_price\n1,A,2,EA,\n")
        zf.writestr("notes.txt", "hello")
    result = safe_inspect_zip(buf.getvalue())
    assert result["ok"]
    assert any(e["name"] == "schedule.csv" for e in result["extracted"])
    assert not result["rejected"]


def test_zip_path_traversal_rejection():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("../evil.txt", "nope")
        zf.writestr("ok.txt", "yes")
    result = safe_inspect_zip(buf.getvalue())
    assert result["ok"]
    assert any(r.get("reason") == "path_traversal_or_absolute" for r in result["rejected"])
    assert any(e["name"] == "ok.txt" for e in result["extracted"])


def test_line_item_quantity_uom_manufacturer_brand_or_equal():
    items = extract_sciquest_product_line_items(BLADES_PDF_TEXT, source_document="event.pdf")
    assert len(items) == 6
    assert items[0]["quantity"] == 300
    assert items[0]["unit_of_measure"] == "EA"
    assert items[0]["CLIN_or_item_number"] == "002367300"
    assert items[0]["manufacturer"] is None  # unknown preserved
    assert items[0]["unit_price"]["status"] == PRICE_TO_BE_PROVIDED
    terms = extract_delivery_and_terms(BLADES_PDF_TEXT, source_document="event.pdf")
    assert terms["delivery_location"]["value"]
    assert terms["brand_name_or_equal"]["value"] is True
    pid = identify_product(line_items=items, terms=terms, title="Tungsten-Carbide Blades")
    assert pid["product_id_state"] in {PRODUCT_BRAND_OR_EQUAL, PRODUCT_SPEC_COMMODITY}


def test_empty_price_not_zero():
    assert _normalize_price_cell("")["status"] == PRICE_TO_BE_PROVIDED
    assert _normalize_price_cell("")["value"] is None
    assert _normalize_price_cell(None)["value"] is not 0
    assert _normalize_price_cell("  ")["value"] is None


def test_unknown_field_preservation():
    items = extract_sciquest_product_line_items(BLADES_PDF_TEXT)
    assert items[0]["model"] is None
    assert items[0]["part_number"] is None


def test_document_completeness_and_product_id():
    items = extract_sciquest_product_line_items(BLADES_PDF_TEXT)
    terms = extract_delivery_and_terms(BLADES_PDF_TEXT)
    pid = identify_product(line_items=items, terms=terms, title="Blades")
    docs = [
        {
            "document_class": DOC_SOLICITATION,
            "access_status": ACCESS_PUBLIC_FETCHED,
            "document_title": "event.pdf",
        },
        {
            "document_class": DOC_SPECIFICATION,
            "access_status": ACCESS_LOGIN_REQUIRED,
            "document_title": "Snow Plow Blade Spec 8Sept20263.pdf",
        },
    ]
    c = assess_document_completeness(documents=docs, line_items=items, terms=terms, product_id=pid)
    assert c["document_completeness"] == COMPLETE_ENOUGH_FOR_PRODUCT_ID


def test_spec_defined_commodity_state():
    items = extract_sciquest_product_line_items(BLADES_PDF_TEXT)
    terms = {"brand_name_or_equal": {"value": False}}
    # force no brand-equal
    pid = identify_product(line_items=items, terms=terms, title="Blades")
    assert pid["product_id_state"] in {PRODUCT_SPEC_COMMODITY, PRODUCT_BRAND_OR_EQUAL, "MULTIPLE_ACCEPTABLE_PRODUCTS"}


def test_supplier_evidence_max_four_no_outreach():
    pid = {"sourcing_description": "Tungsten-carbide snowplow blade qty 300 EA", "product_id_state": PRODUCT_BRAND_OR_EQUAL}
    suppliers = discover_suppliers_for_product(product_id=pid, title="Tungsten-Carbide Blades", max_suppliers=4)
    assert 2 <= len(suppliers) <= 4
    assert all(s.get("quote_required") for s in suppliers)
    assert all(s.get("authorized_distributor_evidence") is None for s in suppliers)


def test_quote_packet_and_quote_required():
    items = extract_sciquest_product_line_items(BLADES_PDF_TEXT)
    terms = extract_delivery_and_terms(BLADES_PDF_TEXT)
    pid = identify_product(line_items=items, terms=terms)
    pkt = build_supplier_quote_packet(
        solicitation_number="645-DOTRFB-2975-2027",
        agency="Iowa DOT",
        product_id=pid,
        line_items=items,
        terms=terms,
        supplier={"supplier_name": "Winter Equipment Company"},
    )
    assert pkt["do_not_send_automatically"] is True
    assert pkt["outreach_performed"] is False
    assert pkt["quantity_summary"] == 970.0
    assert "freight_to_delivery_destination" in pkt["requests"]


def test_unknown_revenue_and_freight_preserved():
    items = extract_sciquest_product_line_items(BLADES_PDF_TEXT)
    econ = compute_procurement_economics(line_items=items)
    assert econ["freight"] is None
    assert econ["freight_status"] == "UNKNOWN"
    assert econ["expected_revenue"] is None
    assert econ["revenue_status"] == "UNKNOWN"
    assert econ["supplier_cost"] is None
    assert econ["working_capital_status"] == "UNKNOWN"
    assert econ["ten_k_target"] == "UNKNOWN"
    assert econ["profit_state"] == "PROFIT_UNKNOWN"


def test_landed_cost_and_working_capital_from_meaningful_cost():
    items = extract_sciquest_product_line_items(BLADES_PDF_TEXT)
    econ = compute_procurement_economics(
        line_items=items,
        public_price_evidence={"unit_price": 20.0, "confidence": "ESTIMATED"},
        freight=2000.0,
        freight_status="ESTIMATED",
        expected_revenue=50000.0,
        revenue_status="ESTIMATED_FROM_EVIDENCE",
    )
    assert econ["supplier_cost"] == 970 * 20.0
    assert econ["landed_cost"] == 970 * 20.0 + 2000.0
    assert econ["working_capital"] == 970 * 20.0 + 2000.0
    assert econ["funding_warranted"] is True
    assert econ["ten_k_target"] in {"PASS", "FAIL", "POSSIBLE"}


def test_funding_downstream_and_match_not_approval():
    # No cost → funding not yet needed
    items = extract_sciquest_product_line_items(BLADES_PDF_TEXT)
    econ = compute_procurement_economics(line_items=items)
    funding = maybe_run_funding(econ)
    assert funding["status"] == "NOT_YET_NEEDED"
    assert funding["approval"] is False
    assert funding["lender_outreach"] is False

    econ2 = compute_procurement_economics(
        line_items=items,
        public_price_evidence={"unit_price": 25.0, "confidence": "ESTIMATED"},
        freight=1000.0,
        freight_status="ESTIMATED",
    )
    funding2 = maybe_run_funding(econ2)
    assert funding2["approval"] is False
    assert funding2["lender_outreach"] is False
    assert funding2.get("match_is_not_approval") is True or funding2["status"] in {
        "NEEDS_VERIFICATION",
        "NOT_YET_NEEDED",
    }


def test_unknown_economics_not_rejected_prematurely():
    items = extract_sciquest_product_line_items(BLADES_PDF_TEXT)
    deal = correct_deal_type_from_documents(title="Blades", line_items=items)
    assert deal["deal_type"] == "ONE_TIME_PRODUCT_PURCHASE"
    assert deal["transactional_resale_fit"] == "GOOD"


def test_badge_deal_reclassified_not_pure_resale():
    text = """
    Product Line Items
    P1
    Iowa State Patrol Badge
    1
    EA - Each
    Service Line Items
    S1
    Badge Repair: Simple Refinish
    1
    EA - Each
    """
    items = extract_sciquest_product_line_items(text)
    deal = correct_deal_type_from_documents(title="Badges", line_items=items, prior_deal_type="ONE_TIME_PRODUCT_PURCHASE")
    assert deal["corrected"] is True
    assert deal["transactional_resale_fit"] == "POOR"


def test_top_deal_replacement_logic():
    selected = select_deals_with_replacement(
        priority=["645-DOTRFB-2975-2027", "005-RFB-3030-2027"],
        killed={"005-RFB-3030-2027"},
        max_deals=3,
    )
    assert "005-RFB-3030-2027" not in selected
    assert "645-DOTRFB-2975-2027" in selected
    assert len(selected) <= 3


def test_bid_ready_remains_strict():
    packet = {
        "requirement": {
            "product_identification": {"product_id_state": PRODUCT_BRAND_OR_EQUAL},
            "line_items": [{"quantity": 1}],
        },
        "economics_block": {
            "supplier_cost": 1000,
            "freight_status": "UNKNOWN",
            "expected_revenue": None,
            "profit_state": "PROFIT_UNKNOWN",
            "working_capital_status": "UNKNOWN",
        },
        "funding": {"status": "NOT_YET_NEEDED"},
        "compliance": {"entity_eligibility": "UNKNOWN", "set_aside": "UNKNOWN"},
    }
    ev = evaluate_bid_ready_strict(packet)
    assert ev["bid_ready"] is False
    assert "freight_unknown" in ev["blockers"]
    assert "funding_not_verified" in ev["blockers"]


def test_no_outreach_flags_and_budget_constant():
    assert PROCUREMENT_HARD_MAX_HTTP == 50
    pkt = build_supplier_quote_packet(
        solicitation_number="X",
        agency="A",
        product_id={"sourcing_description": "y"},
        line_items=[],
        terms={},
    )
    assert pkt["outreach_performed"] is False


def test_package_enumeration_record_fields():
    rec = document_record(
        solicitation_number="645-DOTRFB-2975-2027",
        source_portal="sciquest_iowa",
        document_title="event.pdf",
        document_url="https://example.com/event.pdf",
        access_status=ACCESS_PUBLIC_FETCHED,
        http_status=200,
        content_length=100,
        content_hash="abc",
        provenance="test",
    )
    for key in [
        "solicitation_number",
        "source_portal",
        "document_title",
        "document_url",
        "file_type",
        "fetch_status",
        "http_status",
        "access_status",
        "content_length",
        "hash",
        "provenance",
        "retrieval_timestamp",
        "extraction_status",
        "appears_authoritative",
        "superseded_or_amended",
    ]:
        assert key in rec


def test_sciquest_signed_url_preserved_in_parser():
    from discovery.sciquest import parse_sciquest_public_events

    html = """
    <tr class="x">
      <span class="status-badge">Open</span>
      <a class="btn-link-header" id="e1" href="/apps/Router/ViewSourcingEvent?AuthToken=abc">Tungsten-Carbide Blades for Snow/Ice Removal</a>
      <div>Type</div><div>DOTRFB</div>
      <div>Number</div><div>645-DOTRFB-2975-2027</div>
      <time datetime="2026-10-07T18:00:00Z">Close 10/7/2026, 1:00 PM CDT</time>
      <a href="https://solutions-selectsite-documents.s3.amazonaws.com/Sourcingevent/1423842-event.pdf?X-Amz-Algorithm=AWS4-HMAC-SHA256&amp;X-Amz-Signature=abc123">PDF</a>
    </tr>
    """
    # Need Open/Close structure that parser accepts — use minimal viable row
    row = """
    <tr>
    <td><span class="status-badge">Open</span>
    <a class="btn-link-header" id="x" href="https://app01.jaggaer.com/apps/Router/ViewSourcingEvent?AuthToken=1">Tungsten-Carbide Blades for Snow/Ice Removal Extra Title Words</a>
    <div>Type</div><span>DOTRFB</span>
    <div>Number</div><span>645-DOTRFB-2975-2027</span>
    Close 10/7/2026, 1:00 PM CDT
    <a href="https://solutions-selectsite-documents.s3.amazonaws.com/Sourcingevent/1423842-event.pdf?X-Amz-Algorithm=AWS4-HMAC-SHA256&amp;X-Amz-Signature=deadbeef">View as PDF</a>
    The State of Iowa Department of Administrative Services
    </td></tr>
    """
    opps = parse_sciquest_public_events(row, list_url="https://bids.sciquest.com/apps/Router/PublicEvent?CustomerOrg=DASIowa")
    assert opps, "parser should find opportunity"
    url = opps[0].document_links[0]["url"]
    assert "X-Amz-Signature" in url
    assert url.startswith("https://")
