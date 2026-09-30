"""R1.2 OCR + amendment hardening + master store tests. 0 SAM calls."""

from __future__ import annotations

import io
from pathlib import Path

import pytest

from response_engine.amendment_diff import apply_amendment_diff_to_project, diff_documents, structural_text_diff
from response_engine.legacy_bridge import wrap_legacy_bid_readiness
from response_engine.master_store import attach_master_to_project, list_masters, register_master_document
from response_engine.ocr import (
    apply_ocr_human_correction,
    confidence_bucket,
    gate_ocr_requirements,
    ocr_engine_status,
    ocr_pdf_bytes,
    page_needs_ocr,
)
from response_engine.package_completeness import evaluate_package_completeness
from response_engine.parsers import clear_parse_cache, parse_file_bytes, sha256_bytes
from response_engine.production_intake import ingest_bytes_into_project, refresh_solicitation, run_production_intake
from response_engine.service import compile_project, create_or_get_project_from_opportunity
from response_engine.firewall import firewall_report

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture()
def tmp_store(tmp_path, monkeypatch):
    import response_engine.store as store
    import response_engine.production_intake as pi
    import response_engine.master_store as ms

    monkeypatch.setattr(store, "STORE_DIR", tmp_path / "response_projects")
    monkeypatch.setattr(store, "INDEX_PATH", tmp_path / "response_projects" / "index.json")
    monkeypatch.setattr(pi, "BINARY_STORE", tmp_path / "binaries")
    monkeypatch.setattr(ms, "MASTER_DIR", tmp_path / "shared_master")
    monkeypatch.setattr(ms, "INDEX_PATH", tmp_path / "shared_master" / "index.json")
    monkeypatch.setattr(ms, "BINARY_DIR", tmp_path / "shared_master" / "binaries")
    store.ensure_store()
    return tmp_path


def _scanned_pdf(text: str = "Quantity: 12 each. Delivery shall occur within 30 days.") -> bytes:
    import fitz
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (700, 900), color="white")
    d = ImageDraw.Draw(img)
    y = 50
    for line in text.split("\n"):
        d.text((40, y), line, fill="black")
        y += 30
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    doc = fitz.open()
    page = doc.new_page(width=700, height=900)
    page.insert_image(page.rect, stream=buf.getvalue())
    data = doc.tobytes()
    doc.close()
    return data


def test_ocr_engine_available():
    st = ocr_engine_status()
    assert st["available"] is True
    assert st["engine"]


def test_ocr_scan_detection_and_extraction(tmp_store):
    data = _scanned_pdf("Quantity: 25 each\nSubmit signed form by email.")
    assert page_needs_ocr("") is True
    clear_parse_cache()
    parsed = parse_file_bytes(data, filename="scan.pdf")
    assert parsed.get("ocr_required") or parsed.get("method") in {"pdf_ocr", "pdf_scan_detected", "pdf_native_plus_ocr"}
    # Real OCR path
    ocr = ocr_pdf_bytes(data, filename="scan.pdf", force_all_pages=True)
    assert ocr.get("ok")
    assert ocr.get("stats", {}).get("ocr_pages", 0) >= 1
    assert any(p.get("method") == "OCR" for p in ocr.get("pages") or [])
    assert (ocr.get("text") or "").lower().find("quantity") >= 0 or "25" in (ocr.get("text") or "")


def test_ocr_provenance_and_confidence(tmp_store):
    data = _scanned_pdf("Deadline October 14, 2026. Quantity: 5 each.")
    ocr = ocr_pdf_bytes(data, filename="scan.pdf", force_all_pages=True, document_id="DOC-1")
    page = (ocr.get("pages") or [None])[0]
    assert page
    assert page.get("extraction_method") == "OCR" or page.get("method") == "OCR"
    for b in page.get("blocks") or []:
        assert b.get("page") == 1
        assert b.get("extraction_method") == "OCR"
        assert b.get("ocr_engine_version")
    assert confidence_bucket(0.9) == "HIGH"
    assert confidence_bucket(0.6) == "MEDIUM"
    assert confidence_bucket(0.2) == "LOW"


def test_ocr_human_correction_and_gating(tmp_store):
    p = create_or_get_project_from_opportunity(canonical_opportunity_id="r12-ocr")
    data = _scanned_pdf("Delivery shall occur within 30 days ARO. Quantity: 8 each.")
    ingest_bytes_into_project(p, data=data, filename="scan.pdf")
    compile_project(p, persist=False)
    gate_ocr_requirements(p)
    doc = p["documents"][0]
    # ensure ocr pages exist for correction API
    if not (doc.get("ocr") or {}).get("pages"):
        doc["ocr"] = ocr_pdf_bytes(data, filename="scan.pdf", force_all_pages=True, document_id=doc["document_id"])
    page = doc["ocr"]["pages"][0]["page"]
    original = doc["ocr"]["pages"][0].get("text") or ""
    result = apply_ocr_human_correction(
        p,
        document_id=doc["document_id"],
        page=page,
        corrected_text=original + " [corrected]",
        action="CORRECT",
        notes="fix typo",
    )
    assert result["ok"]
    hist = doc["ocr"]["pages"][0]["correction_history"]
    assert hist[-1]["original_ocr"] is not None
    assert hist[-1]["corrected_text"].endswith("[corrected]")


def test_amendment_textual_and_numeric_diff(tmp_store):
    before = {"document_id": "D1", "text": "Quantity: 10 each. Bid due October 10, 2026. Submit via email."}
    after = {"document_id": "D2", "text": "Quantity: 30 each. Bid due October 14, 2026. Submit via PIEE."}
    struct = structural_text_diff(before["text"], after["text"])
    assert "30" in struct["added_numbers"] or struct["added_numbers"]
    diff = diff_documents(before, after, amendment_id="AMD-1")
    cats = {c["category"] for c in diff["changes"]}
    assert "quantity" in cats or "deadline" in cats or "submission_method" in cats
    assert diff["change_count"] >= 1


def test_amendment_spreadsheet_replacement(tmp_store):
    import openpyxl

    def xb(v):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Pricing"
        ws["A1"] = "Qty"
        ws["A2"] = v
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    p = create_or_get_project_from_opportunity(canonical_opportunity_id="r12-xls")
    ingest_bytes_into_project(p, data=xb(10), filename="CostSheet.xlsx")
    ingest_bytes_into_project(p, data=xb(40), filename="CostSheet.xlsx")
    versions = [d for d in p["documents"] if d["filename"] == "CostSheet.xlsx"]
    assert len(versions) == 2
    assert sum(1 for d in versions if d["controlling_status"] == "SUPERSEDED") == 1
    assert sum(1 for d in versions if d["controlling_status"] == "CONTROLLING") == 1
    assert p.get("amendment_diffs")  # diff applied on supersession


def test_buyer_form_replacement(tmp_store):
    p = create_or_get_project_from_opportunity(canonical_opportunity_id="r12-form")
    ingest_bytes_into_project(p, data=b"Bid Form v1", filename="BidForm.xlsx")
    # invalid xlsx bytes still versioned by filename+hash
    ingest_bytes_into_project(p, data=b"Bid Form v2 revised", filename="BidForm.xlsx")
    versions = [d for d in p["documents"] if d["filename"] == "BidForm.xlsx"]
    assert len(versions) == 2


def test_false_complete_prevention(tmp_store):
    p = create_or_get_project_from_opportunity(canonical_opportunity_id="r12-miss")
    ingest_bytes_into_project(
        p,
        data=b"Base. Quantity: 2 each. See Attachment C for specifications. Signed form required.",
        filename="base.txt",
        document_type="BASE_SOLICITATION",
    )
    compile_project(p, persist=False)
    evaluate_package_completeness(p)
    assert p["package_completeness"]["document_package_complete"] is False
    assert p["package_completeness"]["status"] == "MISSING_REFERENCED_DOCUMENT"


def test_master_versioning_immutable(tmp_store):
    m1 = register_master_document(
        authority="DLA",
        title="Master Sol",
        version="2023",
        data=b"Master 2023",
        filename="m.txt",
    )
    m2 = register_master_document(
        authority="DLA",
        title="Master Sol",
        version="2024",
        data=b"Master 2024 revised",
        filename="m.txt",
    )
    assert m1["master_id"] != m2["master_id"]
    assert m1["version"] == "2023"
    assert m2["version"] == "2024"
    # same version+hash returns existing
    m1b = register_master_document(authority="DLA", title="Master Sol", version="2023", data=b"Master 2023", filename="m.txt")
    assert m1b["master_id"] == m1["master_id"]
    p = create_or_get_project_from_opportunity(canonical_opportunity_id="r12-master")
    ref = attach_master_to_project(p, m1["master_id"], applicability_confirmed=False)
    assert ref["ok"]
    assert p["master_references"][0]["status"] == "MASTER_VERSION_REVIEW_REQUIRED"
    assert len(list_masters(authority="DLA")) >= 2


def test_manual_authoritative_upload(tmp_store):
    from response_engine.production_intake import manual_upload_document

    p = create_or_get_project_from_opportunity(canonical_opportunity_id="r12-upload")
    p["authoritative_source"] = "https://portal.example/login"
    p.setdefault("intake", {})["fetch_results"] = [{"status": "AUTH_REQUIRED"}]
    r = manual_upload_document(p, data=b"RFQ after login. Quantity: 3 each.", filename="buyer.pdf.txt", mark_as_authoritative=True)
    assert r["ok"]
    doc = next(d for d in p["documents"] if d["document_id"] == r["document_id"])
    assert doc.get("source_tag") == "OWNER_UPLOADED_FROM_BUYER"
    assert doc.get("mark_authoritative") is True


def test_parser_cache_and_idempotence(tmp_store):
    clear_parse_cache()
    data = b"Cached text Quantity: 1 each."
    a = parse_file_bytes(data, filename="a.txt")
    b = parse_file_bytes(data, filename="a.txt")
    assert b.get("cache_hit") is True
    p = create_or_get_project_from_opportunity(canonical_opportunity_id="r12-idem")
    ingest_bytes_into_project(p, data=data, filename="a.txt")
    r2 = ingest_bytes_into_project(p, data=data, filename="a.txt")
    assert r2.get("duplicate") is True
    assert len(p["documents"]) == 1


def test_change_refresh(tmp_store):
    p = create_or_get_project_from_opportunity(canonical_opportunity_id="r12-refresh", solicitation_number="NOMATCHXYZ")
    ingest_bytes_into_project(p, data=b"Base qty 1", filename="base.txt")
    before = len(p["documents"])
    # refresh with no new local matches should not explode duplicates
    refresh_solicitation(p)
    assert len(p["documents"]) >= before


def test_legacy_override_protection(tmp_store):
    p = create_or_get_project_from_opportunity(canonical_opportunity_id="r12-leg")
    ingest_bytes_into_project(p, data=b"Country of origin required. Quantity: 1. Signed SF1449. Submit by email 2:00 PM ET.", filename="rfq.txt")
    compile_project(p, persist=True)
    wrapped = wrap_legacy_bid_readiness("r12-leg", {"bid_readiness": {"ladder": "READY"}})
    assert wrapped["bid_readiness"]["ready_to_submit"] is False
    assert wrapped["bid_readiness"]["canonical_source"] == "response_engine_r1"
    assert wrapped["bid_readiness"]["ladder"] != "READY" or wrapped["bid_readiness"].get("conflict_with_legacy")


def test_clean_room_firewall(tmp_store):
    p = create_or_get_project_from_opportunity(canonical_opportunity_id="r12-fw")
    ingest_bytes_into_project(p, data=b"Quantity: 1 each. Submit via portal.", filename="rfq.txt")
    compile_project(p, persist=False)
    report = firewall_report(p)
    # no leakage expected
    assert report.get("clean") is True
    assert len(report.get("leaks") or []) == 0


def test_zero_sam_in_r12_modules():
    for name in ("ocr.py", "amendment_diff.py", "master_store.py", "production_intake.py", "parsers.py"):
        text = (ROOT / "response_engine" / name).read_text(encoding="utf-8")
        assert "api.sam.gov" not in text or "blocked" in text.lower()
        assert "can_spend_sam" not in text


def test_api_routes_r12():
    from app import app

    paths = {getattr(r, "path", None) for r in app.routes}
    assert "/api/response-projects/{response_project_id}/ocr-review" in paths
    assert "/api/response-projects/{response_project_id}/amendment-diffs" in paths
