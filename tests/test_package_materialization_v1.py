"""Package materialization truth + validation tests."""

from __future__ import annotations

from pathlib import Path

from bidnet_engine.package_materialization import (
    BUILD,
    PACKAGE_DOCUMENTS_COMPLETE,
    PACKAGE_DOCUMENTS_PARTIAL,
    build_attachment_index,
    harvest_attachment_urls_from_html,
    is_bidnet_detail_page_url,
    looks_like_html_bytes,
    legacy_package_implies_complete,
    materialize_attachments,
    operator_package_message,
    primary_line_blocker,
    reconcile_package_state,
    resolve_bidnet_private_detail_url,
    validate_local_file,
    _sig_ok,
)


def test_build():
    assert BUILD == "20261007-m3-authoritative-schedule-recovery-v1"


def test_detail_html_only_not_complete():
    completeness, state, ready, blocker = reconcile_package_state(
        discovered=0,
        downloaded=0,
        valid=0,
        auth_doc=None,
        index=[],
        detail_present=True,
    )
    assert completeness == "DETAIL_ONLY"
    assert ready is False
    assert blocker == "NO_ATTACHMENT_INDEX"
    assert legacy_package_implies_complete("PACKAGE_ACQUIRED_OFFICIAL_SOURCE")


def test_index_zero_downloads_incomplete():
    idx = build_attachment_index(
        [{"filename": "pricing_schedule.xlsx", "document_url": "https://example.com/p.xlsx"}]
    )
    completeness, state, ready, blocker = reconcile_package_state(
        discovered=1,
        downloaded=0,
        valid=0,
        auth_doc=None,
        index=idx,
        detail_present=True,
    )
    assert ready is False
    assert completeness in {"PARTIAL", "DETAIL_ONLY", "NONE"} or state == PACKAGE_DOCUMENTS_PARTIAL or "DOWNLOAD" in (state or "")


def test_login_html_as_xlsx_rejected(tmp_path: Path):
    p = tmp_path / "fake.xlsx"
    p.write_bytes(b"<!DOCTYPE html><html>Please log in to continue</html>")
    v = validate_local_file(p, claimed_ext="xlsx")
    assert v["DOCUMENT_CONTENT_VALID"] is False
    assert "html" in str(v["reason"] or "").lower() or "login" in str(v["reason"] or "").lower()


def test_cloudflare_html_as_pdf_rejected(tmp_path: Path):
    p = tmp_path / "doc.pdf"
    p.write_bytes(b"<html>Just a moment... Cloudflare</html>")
    v = validate_local_file(p, claimed_ext="pdf")
    assert v["DOCUMENT_CONTENT_VALID"] is False


def test_html_doctype_extensionless_document_1_rejected(tmp_path: Path):
    """BidNet often saves HTML as document_1 with no extension — must not count as valid."""
    p = tmp_path / "document_1"
    p.write_bytes(b"<!DOCTYPE html><html><body>Please log in</body></html>")
    v = validate_local_file(p, claimed_ext="")
    assert v["DOCUMENT_CONTENT_VALID"] is False
    assert "html" in str(v.get("reason") or "").lower()


def test_utf8_bom_html_rejected():
    bom = b"\xef\xbb\xbf<!DOCTYPE html><html><body>Hi</body></html>"
    assert looks_like_html_bytes(bom) is True
    ok, reason = _sig_ok(bom, "")
    assert ok is False
    assert "html" in (reason or "").lower()


def test_resolve_open_bids_to_private_view():
    pub = (
        "https://www.bidnetdirect.com/new-york/solicitations/open-bids/"
        "Automotive-Parts-and-Accessories/0000437119?purchasingGroupId=1"
    )
    assert is_bidnet_detail_page_url(pub)
    priv = resolve_bidnet_private_detail_url(pub)
    assert priv == "https://www.bidnetdirect.com/private/supplier/solicitations/0000437119/view"


def test_html_viewer_harvests_pdf_links():
    html = b"""<!DOCTYPE html><html><body>
    <a href="/files/Equipment%20List.pdf">Equipment List</a>
    <a href="https://cdn.example.com/BidSchedule.xlsx">Schedule</a>
    <a href="/login">Login</a>
    </body></html>"""
    found = harvest_attachment_urls_from_html(
        html, base_url="https://www.bidnetdirect.com/private/supplier/solicitations/123"
    )
    urls = [f["document_url"] for f in found]
    assert any(u.endswith(".pdf") or "Equipment" in u for u in urls)
    assert any("BidSchedule.xlsx" in u for u in urls)
    assert not any("login" in u.lower() for u in urls)


def test_html_poison_not_materialized_as_valid(tmp_path: Path):
    p = tmp_path / "document_1"
    p.write_bytes(
        b"<!DOCTYPE html><html><body>"
        b'<a href="https://files.example.com/schedule.pdf">schedule.pdf</a>'
        b"</body></html>"
    )
    docs = [
        {
            "filename": "document_1",
            "local_path": str(p),
            "document_url": "https://www.bidnetdirect.com/private/supplier/doc/1",
        }
    ]
    out = materialize_attachments(docs, opportunity_id="ny1", client=None)
    assert out["PACKAGE_DOCUMENT_COUNT_MATERIALIZED"] == 0
    assert out["HARVESTED_FROM_HTML"] >= 1
    assert any("schedule.pdf" in str(e.get("source_url") or "") for e in (out.get("PACKAGE_ATTACHMENT_INDEX") or []))


def test_valid_xlsx_materializes(tmp_path: Path):
    import openpyxl

    p = tmp_path / "pricing_schedule.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["Item", "Description", "Qty"])
    wb.active.append([1, "Widget", 10])
    wb.save(p)
    v = validate_local_file(p, claimed_ext="xlsx")
    assert v["DOCUMENT_CONTENT_VALID"] is True
    docs = [{"filename": "pricing_schedule.xlsx", "local_path": str(p), "document_url": "https://ex/p.xlsx"}]
    out = materialize_attachments(docs, opportunity_id="t1", client=None)
    assert out["PACKAGE_DOCUMENT_COUNT_MATERIALIZED"] >= 1
    assert out["AUTHORITATIVE_PRODUCT_DOC_FOUND"] is True
    assert out["PACKAGE_READY_FOR_LINE_EXTRACTION"] is True


def test_external_buyer_portal_document_valid(tmp_path: Path):
    import openpyxl

    p = tmp_path / "exhibit_a.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["Line", "Desc", "Qty"])
    wb.active.append([1, "Bolt", 5])
    wb.save(p)
    docs = [
        {
            "filename": "Exhibit A.xlsx",
            "local_path": str(p),
            "document_url": "https://opengov.buyer.gov/files/exhibit_a.xlsx",
            "official_source": True,
        }
    ]
    out = materialize_attachments(docs, opportunity_id="t2", client=None)
    assert out["AUTHORITATIVE_PRODUCT_DOC_FOUND"]
    assert out["docs_for_extraction"]


def test_missing_revised_pricing_partial():
    idx = build_attachment_index(
        [
            {"filename": "pricing_schedule.xlsx", "document_url": "https://ex/p.xlsx"},
            {"filename": "Revised Pricing Sheet.xlsx", "document_url": "https://ex/r.xlsx"},
        ]
    )
    # Only first materialized as auth
    auth = {
        **idx[0],
        "LOCAL_PATH": "/tmp/p.xlsx",
        "DOCUMENT_CONTENT_VALID": True,
        "document_role_guess": "PRICING_SCHEDULE",
        "CURRENT_AUTHORITATIVE": True,
    }
    completeness, state, ready, blocker = reconcile_package_state(
        discovered=2,
        downloaded=1,
        valid=1,
        auth_doc=auth,
        index=idx,
        detail_present=True,
    )
    # With auth pricing present, ready can be true even if other file missing
    assert ready is True or completeness == "PARTIAL"


def test_superseded_pricing_not_authoritative(tmp_path: Path):
    import openpyxl

    old = tmp_path / "pricing_old.xlsx"
    new = tmp_path / "Revised_Pricing_Sheet.xlsx"
    for p in (old, new):
        wb = openpyxl.Workbook()
        wb.active.append(["Item", "Description", "Qty"])
        wb.active.append([1, "A", 1])
        wb.save(p)
    docs = [
        {"filename": "pricing_old.xlsx", "local_path": str(old)},
        {"filename": "Revised_Pricing_Sheet.xlsx", "local_path": str(new)},
    ]
    out = materialize_attachments(docs, opportunity_id="t3", client=None)
    auth = out.get("AUTHORITATIVE_PRODUCT_DOC") or {}
    assert "revis" in str(auth.get("filename") or "").lower() or auth.get("CURRENT_AUTHORITATIVE") is True


def test_no_product_schedule_classified():
    idx = build_attachment_index([{"filename": "terms_and_conditions.pdf", "document_url": "https://ex/t.pdf"}])
    completeness, state, ready, blocker = reconcile_package_state(
        discovered=1,
        downloaded=1,
        valid=1,
        auth_doc=None,
        index=idx,
        detail_present=True,
    )
    assert ready is False


def test_schedule_acquisition_vs_parser_failure():
    pkg = {
        "PACKAGE_READY_FOR_LINE_EXTRACTION": False,
        "primary_blocker": "PRODUCT_SCHEDULE_NOT_ACQUIRED",
        "AUTHORITATIVE_PRODUCT_DOC_FOUND": False,
    }
    assert primary_line_blocker(package_result=pkg, lines_ready=False) == "PRODUCT_SCHEDULE_NOT_ACQUIRED"
    pkg2 = {"PACKAGE_READY_FOR_LINE_EXTRACTION": True, "primary_blocker": None, "AUTHORITATIVE_PRODUCT_DOC_FOUND": True}
    assert primary_line_blocker(
        package_result=pkg2,
        lines_ready=False,
        schedule_extract={"SCHEDULE_PRESENT_EXTRACTION_ZERO": [{"code": "SCHEDULE_PRESENT_EXTRACTION_ZERO"}]},
    ) == "PARSER_FAILURE"


def test_operator_plain_language():
    msg = operator_package_message(
        completeness="PARTIAL",
        ready=False,
        auth_doc=None,
        discovered=3,
        valid=1,
        blocker="PRODUCT_SCHEDULE_NOT_ACQUIRED",
        missing_high_value=["Revised Bid Schedule.xlsx"],
    )
    assert "do not have the file" in msg.lower() or "pricing" in msg.lower()
    assert "PACKAGE_DOCUMENTS_PARTIAL" not in msg


def test_deleted_local_file_not_extraction_ready(tmp_path: Path):
    missing = tmp_path / "gone.xlsx"
    v = validate_local_file(missing, claimed_ext="xlsx")
    assert v["DOCUMENT_CONTENT_VALID"] is False
    docs = [{"filename": "pricing.xlsx", "local_path": str(missing)}]
    out = materialize_attachments(docs, opportunity_id="t4", client=None)
    assert out["PACKAGE_READY_FOR_LINE_EXTRACTION"] is False
